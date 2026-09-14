# -*- coding: utf-8 -*-
"""账单解析：先看三份文件的结构"""
import sys, csv, io
from openpyxl import load_workbook

BASE = r"D:\亚欧非大环线\账单"

print("=" * 60)
print("1) 微信支付账单.xlsx")
print("=" * 60)
wb = load_workbook(BASE + r"\微信支付账单.xlsx", data_only=True)
for ws in wb.worksheets:
    print(f"sheet: {ws.title}  dims: {ws.dimensions}  rows: {ws.max_row}  cols: {ws.max_column}")
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        print(row)
        if i >= 8:
            break

print()
print("=" * 60)
print("2) 支付宝交易明细.csv")
print("=" * 60)
raw = open(BASE + r"\支付宝交易明细.csv", "rb").read()
print("first bytes:", raw[:80])
for enc in ("utf-8-sig", "gbk", "utf-8"):
    try:
        text = raw.decode(enc)
        print(f"encoding OK: {enc}")
        break
    except Exception as e:
        print(f"{enc} failed: {e}")
for i, line in enumerate(text.splitlines()[:10]):
    print(i, line)

print()
print("=" * 60)
print("3) 银行卡账单.pdf  (password)")
print("=" * 60)
import pymupdf
doc = pymupdf.open(BASE + r"\银行卡账单.pdf")
print("needs_pass:", doc.needs_pass)
if doc.needs_pass:
    doc.authenticate("315714")
print("pages:", doc.page_count)
for pno in range(min(doc.page_count, 2)):
    txt = doc[pno].get_text()
    print(f"--- page {pno+1} ({len(txt)} chars) ---")
    print(txt[:1500])

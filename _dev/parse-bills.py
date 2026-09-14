# -*- coding: utf-8 -*-
"""统一解析三份账单 → 结构化 JSON + 可读明细表"""
import json, re, csv, io
from openpyxl import load_workbook
import pymupdf

BASE = r"D:\亚欧非大环线\账单"
OUT_JSON = r"D:\亚欧非大环线\_dev\bills-structured.json"
OUT_TXT = r"D:\亚欧非大环线\_dev\bills-cleaned.txt"

records = []

# ---------- 1) 微信 xlsx ----------
wb = load_workbook(BASE + r"\微信支付账单.xlsx", data_only=True)
ws = wb["Sheet1"]
rows = list(ws.iter_rows(values_only=True))
# 表头在第 18 行(索引17), 数据从 19 行(索引18) 开始
for r in rows[18:]:
    if r[0] is None:
        continue
    rec = {
        "source": "wechat",
        "datetime": str(r[0]).strip(),
        "type": str(r[1]).strip() if r[1] else "",
        "counterparty": str(r[2]).strip() if r[2] else "",
        "goods": str(r[3]).strip() if r[3] else "",
        "dir": str(r[4]).strip() if r[4] else "",
        "amount": float(r[5]) if r[5] is not None else 0.0,
        "pay_method": str(r[6]).strip() if r[6] else "",
        "status": str(r[7]).strip() if r[7] else "",
        "remark": str(r[10]).strip() if len(r) > 10 and r[10] else "",
    }
    records.append(rec)

# ---------- 2) 支付宝 csv ----------
text = open(BASE + r"\支付宝交易明细.csv", "rb").read().decode("gbk")
lines = text.splitlines()
# 找表头行（第 24 行，索引 23）
hdr_idx = None
for i, ln in enumerate(lines):
    if ln.startswith("交易时间,交易分类"):
        hdr_idx = i
        break
for ln in lines[hdr_idx + 1:]:
    ln = ln.strip().rstrip(",")
    if not ln:
        continue
    parts = next(csv.reader(io.StringIO(ln)))
    parts = [p.strip() for p in parts]
    if len(parts) < 9:
        continue
    t, cat, cparty, acct, goods, d, amt, method, status, oid, mid, remark = (parts + [""] * 12)[:12]
    if not t.strip():
        continue
    records.append({
        "source": "alipay",
        "datetime": t.strip(),
        "type": cat.strip(),
        "counterparty": cparty.strip(),
        "goods": goods.strip(),
        "dir": "收入" if d.strip() == "收入" else ("支出" if d.strip() == "支出" else "不计收支"),
        "amount": float(amt) if amt.strip() else 0.0,
        "pay_method": method.strip(),
        "status": status.strip(),
        "remark": remark.strip(),
    })

# ---------- 3) 工商银行信用卡 PDF ----------
doc = pymupdf.open(BASE + r"\银行卡账单.pdf")
if doc.needs_pass:
    doc.authenticate("315714")
date_re = re.compile(r"^\d{4}-\d{2}-\d{2}$")
cur = None
pdf_records = []
for pno in range(doc.page_count):
    for ln in doc[pno].get_text().splitlines():
        ln = ln.strip()
        if not ln:
            continue
        if date_re.match(ln):
            if cur:
                pdf_records.append(cur)
            cur = {"source": "icbc", "datetime": ln, "time": None, "card": None, "dir": None,
                   "ccy": None, "amount": None, "settle_ccy": None, "settle_amt": None,
                   "balance": None, "summary": None, "merchant": None}
            continue
        if cur is None:
            continue
        if ln.startswith("本页支出算术合计") or ln.startswith("本页交易笔数") or ln.startswith("本页收入算术合计"):
            if cur:
                pdf_records.append(cur)
                cur = None
            continue
        if cur["time"] is None:
            cur["time"] = ln
        elif cur["card"] is None:
            cur["card"] = ln
        elif cur["dir"] is None:
            cur["dir"] = ln
        elif cur["ccy"] is None:
            cur["ccy"] = ln
        elif cur["amount"] is None:
            cur["amount"] = ln
        elif cur["settle_ccy"] is None:
            cur["settle_ccy"] = ln
        elif cur["settle_amt"] is None:
            cur["settle_amt"] = ln
        elif cur["balance"] is None:
            cur["balance"] = ln
        elif cur["summary"] is None:
            cur["summary"] = ln
        else:
            cur["merchant"] = (cur["merchant"] + " " + ln).strip() if cur["merchant"] else ln
if cur:
    pdf_records.append(cur)

def to_num(s):
    if s is None:
        return None
    return float(s.replace(",", "").strip())

for r in pdf_records:
    records.append({
        "source": "icbc",
        "datetime": r["datetime"],
        "type": r["summary"] or "",
        "counterparty": r["merchant"] or "",
        "goods": "",
        "dir": "贷" if r["dir"] == "贷" else "借",
        "amount": to_num(r["amount"]),
        "amount_ccy": r["ccy"],
        "settle_amount": to_num(r["settle_amt"]),
        "settle_ccy": r["settle_ccy"],
        "pay_method": "工行信用卡",
        "status": "",
        "remark": "",
    })

# ---------- 输出 ----------
with open(OUT_JSON, "w", encoding="utf-8") as f:
    json.dump(records, f, ensure_ascii=False, indent=1)

buf = []
for i, r in enumerate(records, 1):
    if r["source"] == "icbc":
        buf.append(f"[{i}] {r['source']} {r['datetime']} {r.get('time') or ''} {r['dir']} "
                   f"{r['amount']} {r['amount_ccy']} -> {r['settle_amount']} {r['settle_ccy']} "
                   f"({r['type']}) {r['counterparty']}")
    else:
        buf.append(f"[{i}] {r['source']} {r['datetime']} {r['dir']} {r['amount']}元 "
                   f"({r['type']}) {r['counterparty']} | {r['goods']} | {r['status']}")

with open(OUT_TXT, "w", encoding="utf-8") as f:
    f.write("\n".join(buf))

# 汇总统计
from collections import Counter
print("总记录数:", len(records))
print("来源:", dict(Counter(r["source"] for r in records)))
wx_exp = sum(r["amount"] for r in records if r["source"] == "wechat" and r["dir"] == "支出")
wx_inc = sum(r["amount"] for r in records if r["source"] == "wechat" and r["dir"] == "收入")
ali_exp = sum(r["amount"] for r in records if r["source"] == "alipay" and r["dir"] == "支出")
ali_neutral = sum(r["amount"] for r in records if r["source"] == "alipay" and r["dir"] == "不计收支")
icbc_borrow = sum(r["settle_amount"] for r in records if r["source"] == "icbc" and r["dir"] == "借")
icbc_refund = sum(r["settle_amount"] for r in records if r["source"] == "icbc" and r["dir"] == "贷" and r["type"] == "退货")
icbc_repay = sum(r["settle_amount"] for r in records if r["source"] == "icbc" and r["dir"] == "贷" and r["type"] == "转帐")
icbc_other_credit = sum(r["settle_amount"] for r in records if r["source"] == "icbc" and r["dir"] == "贷" and r["type"] not in ("退货", "转帐"))
print(f"微信: 支出 {wx_exp:.2f}  收入 {wx_inc:.2f}")
print(f"支付宝: 支出 {ali_exp:.2f}  不计收支 {ali_neutral:.2f}")
print(f"工行卡: 借(支出) {icbc_borrow:.2f} USD  贷-退货 {icbc_refund:.2f} USD  贷-还款 {icbc_repay:.2f} USD  其他贷 {icbc_other_credit:.2f}")
print(f"工行卡净支出 = {icbc_borrow - icbc_refund:.2f} USD (还款总额 {icbc_repay:.2f} USD)")
print("\n明细已写入:", OUT_TXT)

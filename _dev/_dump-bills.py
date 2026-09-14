# -*- coding: utf-8 -*-
"""导出三份账单的完整内容到文本文件，便于逐条分析"""
import pymupdf
from openpyxl import load_workbook

BASE = r"D:\亚欧非大环线\账单"
OUT = r"D:\亚欧非大环线\_dev\_bills-dump.txt"
lines = []

# 1) 微信 xlsx 全部行
wb = load_workbook(BASE + r"\微信支付账单.xlsx", data_only=True)
ws = wb["Sheet1"]
lines.append("===== 微信支付账单（xlsx 全部 86 行）=====")
for i, row in enumerate(ws.iter_rows(values_only=True), 1):
    lines.append(f"{i}\t" + "\t".join("" if c is None else str(c) for c in row))

# 2) 支付宝 csv 全部行（gbk）
lines.append("\n===== 支付宝交易明细（csv 全部行）=====")
text = open(BASE + r"\支付宝交易明细.csv", "rb").read().decode("gbk")
for i, line in enumerate(text.splitlines(), 1):
    lines.append(f"{i}\t{line}")

# 3) 银行卡 PDF 全部文本
lines.append("\n===== 工商银行信用卡账单（PDF 全部 7 页）=====")
doc = pymupdf.open(BASE + r"\银行卡账单.pdf")
if doc.needs_pass:
    doc.authenticate("315714")
for pno in range(doc.page_count):
    lines.append(f"\n----- PDF page {pno + 1} -----")
    lines.append(doc[pno].get_text())

with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("dumped to", OUT, "chars:", sum(len(l) for l in lines))
